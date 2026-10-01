from pathlib import Path
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'docs' / '시야체크_AI_기획안_개정본.md'
OUTPUT = ROOT / 'docs' / '시야체크_AI_기획안_개정본.docx'
import tempfile
ART = Path(tempfile.gettempdir()) / 'sightcheck-proposal'
ART.mkdir(parents=True, exist_ok=True)

def font(run):
    run.font.name = 'Malgun Gothic'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

def runs(p, text):
    # Preserve source links as clickable titles, without printing long URLs.
    for part in re.split(r'(\[[^\]]+\]\(https?://[^\s)]+\))', text):
        match = re.fullmatch(r'\[([^\]]+)\]\((https?://[^\s)]+)\)', part)
        if not match:
            r=p.add_run(part);font(r);continue
        link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(match[2],RT.HYPERLINK,is_external=True))
        r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
        f=OxmlElement('w:rFonts');f.set(qn('w:eastAsia'),'Malgun Gothic');f.set(qn('w:ascii'),'Malgun Gothic');pr.append(f)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),'24');pr.append(sz)
        color=OxmlElement('w:color');color.set(qn('w:val'),'315B77');pr.append(color)
        r.append(pr);t=OxmlElement('w:t');t.text=match[1];r.append(t);link.append(r);p._p.append(link)

doc=Document();s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.6);s.bottom_margin=Cm(1.6);s.left_margin=Cm(1.75);s.right_margin=Cm(1.75)
s.footer_distance=Cm(.65);s.header_distance=Cm(.65)
for grid in s._sectPr.xpath('./w:docGrid'):s._sectPr.remove(grid)
for name,size,bold in [('Normal',12,False),('Title',15,True),('Heading 1',12,True)]:
    st=doc.styles[name];st.font.name='Malgun Gothic';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
    # Korean proposal guidance is 160% of the 12pt font = 19.2pt.
    # Word's font-metric multiple yields much larger lines for Malgun Gothic.
    st.paragraph_format.line_spacing=Pt(size * 1.6);st.paragraph_format.space_after=Pt(4)
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');st._element.get_or_add_pPr().append(snap)
    if name!='Normal':st.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(7)
doc.styles['Title'].paragraph_format.space_after=Pt(10)
for border in doc.styles['Title']._element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)

# Only original synthetic plan and prototype screenshots are used; no third-party photos.
plan=Image.open(ROOT/'public/engine-plan.png').convert('RGB')
compare=Image.open(ROOT/'screenshots/engine-compare.png').convert('RGB')
strip=Image.new('RGB',(1200,350),'white');d=ImageDraw.Draw(strip)
plan.thumbnail((310,290));strip.paste(plan,(0,45))
crop=compare.crop((307,350,939,615));crop.thumbnail((760,290));strip.paste(crop,(410,50))
label=ImageFont.truetype('C:/Windows/Fonts/malgunbd.ttf',24)
d.text((5,5),'입력: 가상 축척 도면',font=label,fill='#28352a');d.text((415,5),'결과: C3와 D3의 시야 비교',font=label,fill='#28352a')
d.line((328,170,382,170),fill='#617757',width=6);d.polygon([(382,170),(366,159),(366,181)],fill='#617757')
strip.save(ART/'proof-strip.png')

active=None
for line in SOURCE.read_text(encoding='utf-8-sig').splitlines():
    line=line.strip()
    if not line:continue
    if line.startswith('# '):
        p=doc.add_paragraph(style='Title');runs(p,line[2:])
    elif line.startswith('## '):
        section=line[3:]
        if section.startswith('4.') and active=='3':
            p=doc.add_paragraph();p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(2)
            p.add_run().add_picture(str(ART/'proof-strip.png'),width=Cm(17))
            p=doc.add_paragraph();runs(p,'가상 축척 도면 → 실제 시제품의 C3·D3 비교. 현장 정확도 검증 자료는 아님.')
        p=doc.add_paragraph(style='Heading 1');runs(p,section)
        if section.startswith(('3.','4.')):p.paragraph_format.page_break_before=True
        active=section.split('.')[0]
    else:
        p=doc.add_paragraph();runs(p,line)
        p.paragraph_format.widow_control=True

footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('시야체크 AI  ·  ');font(r);r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='시야체크 AI 기획안 개정본';doc.core_properties.author='';doc.core_properties.subject='AI 활용 아이디어 공모전 6개 항목'
doc.save(OUTPUT);print(OUTPUT)
